"""
Data Manager for 5 Beijing Yizhuang BDA Communities
Parses '五社区无人配送模型数据汇总.xlsx' and provides clean data structures.
"""

import os
import re
import math
import openpyxl

class DataManager:
    def __init__(self, excel_path=None):
        if excel_path is None:
            # Default path relative to this script or workspace
            base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
            parent_dir = os.path.dirname(base_dir)
            candidates = [
                os.path.join(parent_dir, 'data', '五社区无人配送模型数据汇总.xlsx'),
                os.path.join(base_dir, 'data', '五社区无人配送模型数据汇总.xlsx'),
                os.path.join(parent_dir, 'G数据_已汇总', 'G数据（已汇总）', '五社区无人配送模型数据汇总.xlsx'),
                os.path.join(base_dir, '五社区无人配送模型数据汇总.xlsx')
            ]
            excel_path = next((p for p in candidates if os.path.exists(p)), candidates[0])
        
        self.excel_path = excel_path
        self.communities = {}
        self.demand_nodes = {}
        self.facility_nodes = {}
        self.road_nodes = {}
        self.road_edges = {}
        self.global_params = {}
        self.hub_coord = {"lng": 116.506, "lat": 39.798, "name": "亦庄片区物流综合枢纽分拨中心 (HUB)"}

        if os.path.exists(self.excel_path):
            self.load_all()

    def parse_wkt_polygon(self, wkt_str):
        if not wkt_str:
            return []
        coords = []
        matches = re.findall(r'([0-9\.]+)\s+([0-9\.]+)', str(wkt_str))
        for lng, lat in matches:
            coords.append([float(lat), float(lng)])  # Leaflet uses [lat, lng]
        return coords

    def load_all(self):
        wb = openpyxl.load_workbook(self.excel_path, data_only=True)
        
        # 1. 社区表
        if '社区' in wb.sheetnames:
            ws = wb['社区']
            rows = list(ws.iter_rows(values_only=True))
            # Header is typically at row 4 (0-indexed 4 or 3)
            header_idx = None
            for idx, r in enumerate(rows[:10]):
                if r and '社区ID' in [str(c).strip() for c in r if c is not None]:
                    header_idx = idx
                    break
            if header_idx is not None:
                headers = [str(c).strip() if c is not None else '' for c in rows[header_idx]]
                for r in rows[header_idx+1:]:
                    if not r or not r[0]:
                        continue
                    row_dict = dict(zip(headers, r))
                    cid = str(row_dict.get('社区ID', '')).strip()
                    if cid and cid.startswith('C'):
                        poly = self.parse_wkt_polygon(row_dict.get('边界坐标WKT（WGS84）', ''))
                        self.communities[cid] = {
                            "id": cid,
                            "name": str(row_dict.get('社区名称', '')).strip(),
                            "street": str(row_dict.get('所属社区/街道', '')).strip(),
                            "address": str(row_dict.get('地址', '')).strip(),
                            "area_sqm": float(row_dict.get('占地面积㎡', 0) or 0),
                            "build_area_sqm": float(row_dict.get('建筑面积㎡', 0) or 0),
                            "plot_ratio": float(row_dict.get('容积率', 0) or 0),
                            "building_count": int(row_dict.get('楼栋数', 0) or 0),
                            "households": int(row_dict.get('总户数', 0) or 0),
                            "polygon": poly,
                            "center": self._calculate_center(poly)
                        }

        # 2. 需求节点表 (楼栋)
        if '需求节点' in wb.sheetnames:
            ws = wb['需求节点']
            rows = list(ws.iter_rows(values_only=True))
            header_idx = None
            for idx, r in enumerate(rows[:10]):
                if r and '需求节点ID' in [str(c).strip() for c in r if c is not None]:
                    header_idx = idx
                    break
            if header_idx is not None:
                headers = [str(c).strip() if c is not None else '' for c in rows[header_idx]]
                for r in rows[header_idx+1:]:
                    if not r or not r[0]:
                        continue
                    row_dict = dict(zip(headers, r))
                    nid = str(row_dict.get('需求节点ID', '')).strip()
                    cid = str(row_dict.get('社区ID', '')).strip()
                    if nid:
                        lng = float(row_dict.get('经度WGS84', 0) or 0)
                        lat = float(row_dict.get('纬度WGS84', 0) or 0)
                        self.demand_nodes[nid] = {
                            "id": nid,
                            "community_id": cid,
                            "name": str(row_dict.get('楼栋/分区名称', '')).strip(),
                            "lng": lng,
                            "lat": lat,
                            "building_type": str(row_dict.get('建筑类型', '')).strip(),
                            "households": int(row_dict.get('户数', 0) or 0),
                            "daily_pkgs": float(row_dict.get('日均件量', 0) or 0),
                            "peak_pkgs": float(row_dict.get('峰值件量', 0) or 0),
                            "time_window": str(row_dict.get('服务时间窗', '08:00—21:00')).strip()
                        }

        # 3. 设施点表
        if '设施点' in wb.sheetnames:
            ws = wb['设施点']
            rows = list(ws.iter_rows(values_only=True))
            header_idx = None
            for idx, r in enumerate(rows[:10]):
                if r and '设施ID' in [str(c).strip() for c in r if c is not None]:
                    header_idx = idx
                    break
            if header_idx is not None:
                headers = [str(c).strip() if c is not None else '' for c in rows[header_idx]]
                for r in rows[header_idx+1:]:
                    if not r or not r[0]:
                        continue
                    row_dict = dict(zip(headers, r))
                    fid = str(row_dict.get('设施ID', '')).strip()
                    cid = str(row_dict.get('社区ID', '')).strip()
                    if fid:
                        lng = float(row_dict.get('经度WGS84', 0) or 0)
                        lat = float(row_dict.get('纬度WGS84', 0) or 0)
                        self.facility_nodes[fid] = {
                            "id": fid,
                            "community_id": cid,
                            "name": str(row_dict.get('设施名称', '')).strip(),
                            "type": str(row_dict.get('设施类型', '')).strip(),
                            "status": str(row_dict.get('设施状态', '')).strip(),
                            "lng": lng,
                            "lat": lat,
                            "capacity_str": str(row_dict.get('容量', '')).strip(),
                            "current_util": float(row_dict.get('当前利用量件/日', 0) or 0),
                            "open_time": str(row_dict.get('开放时间', '00:00—24:00')).strip()
                        }

        # 4. 路网节点与路网边
        if '路网节点' in wb.sheetnames:
            ws = wb['路网节点']
            rows = list(ws.iter_rows(values_only=True))
            header_idx = None
            for idx, r in enumerate(rows[:10]):
                if r and '路网节点ID' in [str(c).strip() for c in r if c is not None]:
                    header_idx = idx
                    break
            if header_idx is not None:
                headers = [str(c).strip() if c is not None else '' for c in rows[header_idx]]
                for r in rows[header_idx+1:]:
                    if not r or not r[0]:
                        continue
                    row_dict = dict(zip(headers, r))
                    rn_id = str(row_dict.get('路网节点ID', '')).strip()
                    cid = str(row_dict.get('社区ID', '')).strip()
                    if rn_id:
                        lng = float(row_dict.get('经度WGS84', 0) or 0)
                        lat = float(row_dict.get('纬度WGS84', 0) or 0)
                        self.road_nodes[rn_id] = {
                            "id": rn_id,
                            "community_id": cid,
                            "type": str(row_dict.get('节点类型', '')).strip(),
                            "lng": lng,
                            "lat": lat
                        }

        if '路网边' in wb.sheetnames:
            ws = wb['路网边']
            rows = list(ws.iter_rows(values_only=True))
            header_idx = None
            for idx, r in enumerate(rows[:10]):
                if r and '路段ID' in [str(c).strip() for c in r if c is not None]:
                    header_idx = idx
                    break
            if header_idx is not None:
                headers = [str(c).strip() if c is not None else '' for c in rows[header_idx]]
                for r in rows[header_idx+1:]:
                    if not r or not r[0]:
                        continue
                    row_dict = dict(zip(headers, r))
                    edge_id = str(row_dict.get('路段ID', '')).strip()
                    cid = str(row_dict.get('社区ID', '')).strip()
                    if edge_id:
                        self.road_edges[edge_id] = {
                            "id": edge_id,
                            "community_id": cid,
                            "from_node": str(row_dict.get('起点ID', '')).strip(),
                            "to_node": str(row_dict.get('终点ID', '')).strip(),
                            "length_m": float(row_dict.get('距离m', 0) or 0),
                            "is_bidirectional": str(row_dict.get('是否双向', '是')).strip() == '是',
                            "unmanned_passable": str(row_dict.get('无人车可通行', '是')).strip() == '是',
                            "width_m": float(row_dict.get('道路净宽m', 3.0) or 3.0),
                            "slope_pct": float(row_dict.get('坡度%', 0) or 0),
                            "obstacle": str(row_dict.get('障碍类型', '无')).strip()
                        }

        wb.close()

    def _calculate_center(self, poly):
        if not poly:
            return [39.798, 116.485]
        lats = [p[0] for p in poly]
        lngs = [p[1] for p in poly]
        return [sum(lats)/len(lats), sum(lngs)/len(lngs)]

    def get_community_summary(self, cid):
        c = self.communities.get(cid, {})
        d_nodes = [d for d in self.demand_nodes.values() if d['community_id'] == cid]
        f_nodes = [f for f in self.facility_nodes.values() if f['community_id'] == cid]
        r_edges = [e for e in self.road_edges.values() if e['community_id'] == cid]
        r_nodes = [n for n in self.road_nodes.values() if n['community_id'] == cid]

        # 预组装路网拓扑段坐标，起终点坐标直出，前端零计算损耗
        enriched_edges = []
        for e in r_edges:
            fn = self.road_nodes.get(e['from_node'])
            tn = self.road_nodes.get(e['to_node'])
            edge_item = dict(e)
            if fn and tn:
                edge_item['from_coord'] = [fn['lat'], fn['lng']]
                edge_item['to_coord'] = [tn['lat'], tn['lng']]
            enriched_edges.append(edge_item)
        
        tot_daily = sum(d['daily_pkgs'] for d in d_nodes)
        tot_peak = sum(d['peak_pkgs'] for d in d_nodes)
        tot_households = sum(d['households'] for d in d_nodes) or c.get('households', 0)

        return {
            "info": c,
            "demand_nodes": d_nodes,
            "facilities": f_nodes,
            "road_nodes": r_nodes,
            "road_edges": enriched_edges,
            "road_edge_count": len(r_edges),
            "total_households": tot_households,
            "total_daily_pkgs": tot_daily,
            "total_peak_pkgs": tot_peak
        }

    def get_all_overview(self):
        result = []
        for cid in sorted(self.communities.keys()):
            s = self.get_community_summary(cid)
            result.append({
                "id": cid,
                "name": s["info"].get("name", cid),
                "street": s["info"].get("street", ""),
                "households": s["total_households"],
                "daily_pkgs": s["total_daily_pkgs"],
                "peak_pkgs": s["total_peak_pkgs"],
                "building_nodes": len(s["demand_nodes"]),
                "facilities": len(s["facilities"]),
                "center": s["info"].get("center", [39.798, 116.485]),
                "polygon": s["info"].get("polygon", []),
                "area_sqm": s["info"].get("area_sqm", 0),
                "building_count": s["info"].get("building_count", 0)
            })
        return result

if __name__ == "__main__":
    dm = DataManager()
    print("Loaded communities:", list(dm.communities.keys()))
    overview = dm.get_all_overview()
    for o in overview:
        print(o)

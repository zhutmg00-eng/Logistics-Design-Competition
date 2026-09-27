import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document(r'C:/Users/zhutmg/Desktop/第七章平台(1).docx')
print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

for i, p in enumerate(doc.paragraphs):
    hl_runs = []
    for r in p.runs:
        rPr = r._r.rPr
        if rPr is not None:
            hl = rPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}highlight')
            if hl is not None:
                val = hl.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')
                hl_runs.append((r.text, val))
    
    # check bold runs
    bold_runs = [r.text for r in p.runs if r.bold]
    
    # Print paragraph if interesting
    is_heading = p.text.startswith('7.') or p.text.startswith('7')
    has_fig = '图' in p.text or '【' in p.text
    if hl_runs or is_heading or has_fig or len(p.text.strip()) < 30:
        print(f"P{i:03d} [style={p.style.name}]: {p.text}")
        if hl_runs:
            print(f"     >>> HIGHLIGHT: {hl_runs}")
        if bold_runs:
            print(f"     >>> BOLD: {bold_runs}")

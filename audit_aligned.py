import docx
import zipfile
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

target = r'C:\Users\zhutmg\Desktop\第七章平台(1)_对齐排版与4K高清版.docx'
print("=== 1. File Size ===")
print(f"Size: {os.path.getsize(target)/(1024*1024):.2f} MB ({os.path.getsize(target)} bytes)")

print("\n=== 2. Embedded Media in DOCX ===")
with zipfile.ZipFile(target, 'r') as zf:
    for f in sorted(zf.namelist()):
        if f.startswith('word/media/'):
            print(f"  {f:30s} {zf.getinfo(f).file_size/1024:8.1f} KB")

doc = docx.Document(target)
print(f"\n=== 3. Paragraphs & Tables ===")
print(f"Total paragraphs: {len(doc.paragraphs)}, Tables: {len(doc.tables)}")

print("\n=== 4. Section Properties ===")
for i, sec in enumerate(doc.sections):
    print(f"Sec {i}: w={sec.page_width.pt}pt, h={sec.page_height.pt}pt, top={sec.top_margin.pt}pt, bottom={sec.bottom_margin.pt}pt, left={sec.left_margin.pt}pt, right={sec.right_margin.pt}pt")

print("\n=== 5. Headings and Captions ===")
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t.startswith('7') or t.startswith('图7-') or t.startswith('表7-'):
        r = p.runs[0] if p.runs else None
        f_info = f"{r.font.name} {r.font.size.pt if r.font.size else 'None'}pt bold={r.font.bold}" if r else "no runs"
        print(f"  [P{i:02d}] ({f_info}): {t}")

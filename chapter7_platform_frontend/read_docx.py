import docx
import os

doc_path = os.path.join(os.path.expanduser("~"), "Desktop", "第七章平台.docx")
doc = docx.Document(doc_path)
with open("docx_text.txt", "w", encoding="utf-8") as f:
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if txt:
            f.write(f"P{i}: {txt}\n")
    for t_idx, table in enumerate(doc.tables):
        f.write(f"\n--- TABLE {t_idx} ---\n")
        for row in table.rows:
            row_txt = " | ".join([cell.text.strip().replace('\n', ' ') for cell in row.cells])
            f.write(row_txt + "\n")
print("Done writing docx_text.txt")

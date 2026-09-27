import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document(r'C:/Users/zhutmg/Desktop/第七章平台(1).docx')
with open('d:/物流设计大赛/doc_full_dump.txt', 'w', encoding='utf-8') as f:
    for i, p in enumerate(doc.paragraphs):
        hl_list = []
        for r in p.runs:
            rPr = r._r.rPr
            if rPr is not None:
                hl = rPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}highlight')
                if hl is not None:
                    val = hl.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')
                    hl_list.append((r.text, val))
        
        f.write(f"--- P{i:03d} (Style: {p.style.name}) ---\n")
        if hl_list:
            f.write(f"HIGHLIGHTS: {hl_list}\n")
        f.write(f"{p.text}\n\n")

    for t_idx, table in enumerate(doc.tables):
        f.write(f"=== Table {t_idx} ===\n")
        for r_idx, row in enumerate(table.rows):
            cells = [c.text.replace('\n', ' ') for c in row.cells]
            f.write(f"Row {r_idx}: {' | '.join(cells)}\n")

print("Dumped full document to d:/物流设计大赛/doc_full_dump.txt")

# -*- coding: utf-8 -*-
import os
import sys
import zipfile
import docx

sys.stdout.reconfigure(encoding='utf-8')

target_file = r'C:\Users\zhutmg\Desktop\第七章平台(1).docx'

print("=== 1. Checking File Existence and Size ===")
if not os.path.exists(target_file):
    print("[ERROR] File does not exist!")
    sys.exit(1)

size_bytes = os.path.getsize(target_file)
print(f"[OK] File size: {size_bytes / (1024*1024):.2f} MB ({size_bytes} bytes)")

print("\n=== 2. Checking Embedded Images in DOCX Container ===")
with zipfile.ZipFile(target_file, 'r') as zf:
    media_files = [f for f in zf.namelist() if f.startswith('word/media/')]
    print(f"[OK] Total embedded media files in zip: {len(media_files)}")
    for mf in sorted(media_files):
        info = zf.getinfo(mf)
        print(f"   -> {mf:30s} {info.file_size / 1024:8.1f} KB")

print("\n=== 3. Checking Document Elements with python-docx ===")
doc = docx.Document(target_file)
print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

headings = []
figures = []
tables = []

for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if text.startswith('7') or text.startswith('表7-'):
        headings.append((i, text))
    if text.startswith('图7-'):
        figures.append((i, text))

print("\n=== 4. Headings List ===")
for idx, h in headings:
    print(f"  [P{idx:03d}] {h}")

print(f"\nTotal Headings Found: {len(headings)}")

print("\n=== 5. Figures Captions List ===")
for idx, f in figures:
    print(f"  [P{idx:03d}] {f}")

print(f"\nTotal Figures Found: {len(figures)}")

print("\n=== 6. Table Inspection ===")
for t_idx, table in enumerate(doc.tables):
    print(f"Table {t_idx}: {len(table.rows)} rows x {len(table.columns)} cols")
    for r_idx, row in enumerate(table.rows):
        cells = [c.text.replace('\n', ' ') for c in row.cells]
        print(f"  Row {r_idx}: {' | '.join(cells[:2])} ...")

print("\n=== 7. Searching for Any Leftover Placeholders ===")
bad_keywords = ['【图', '建议表现', '建议采用', '订单进入→服务需求确认', 'C04-R2\n执行车辆']
found_bad = False
for i, p in enumerate(doc.paragraphs):
    for kw in bad_keywords:
        if kw in p.text:
            print(f"[WARNING] Found suspicious placeholder at P{i:03d}: {p.text}")
            found_bad = True

if not found_bad:
    print("[PERFECT] Zero leftover draft placeholders found! All sections cleanly integrated.")

print("\nAuditing complete!")
